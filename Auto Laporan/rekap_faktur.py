"""Rekap jumlah No. Faktur per Nama Penjual dari PDF Laporan Penjualan (ACCURATE).

No. Faktur yang sama (beberapa baris barang) dihitung 1 kali.

Pemakaian:
    python rekap_faktur.py                 -> proses semua PDF di folder ini
    python rekap_faktur.py laporan.pdf     -> proses file tertentu

Hasil: "Rekap Faktur - <nama pdf>.xlsx" di folder yang sama dengan PDF-nya.
"""
import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NO_SELLER = "(Tanpa Nama Penjual)"
# A sales line reads: [Nama Penjual] <No. Faktur> [Nama depan Penjual] <Tgl Faktur> ...
# Some report layouts put extra columns (e.g. "Nama depan Penjual") between the invoice and the date,
# so up to 5 words may sit in between.
DATE = r"(\d{1,2}\s+[A-Za-z]{3,}\s+\d{4})"
GAP = r"\s+(?:\S+\s+){0,5}?"
# Typical numbers (PINV-260925-0038, SI.2026.0012, INV/2609/0031), even when glued to the seller name.
LINE_RE_STRICT = re.compile(r"^(.*?)\s*([A-Z]{2,}[\-/.][A-Za-z0-9/.\-]*\d[A-Za-z0-9/.\-]*)" + GAP + DATE)
# Any other token with letters and digits, directly before the date.
LINE_RE_LOOSE = re.compile(r"^(?:(.*?)\s+)?((?=\S*[A-Za-z])(?=\S*\d)[A-Za-z0-9]\S*)\s+" + DATE)
SKIP_RE = re.compile(r"^(Dari|Total dari|Cetak di)\b", re.I)

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def extract_invoices(pdf_path):
    """Return {(seller, no_faktur): tgl_faktur} in report order, plus the report period."""
    invoices = {}
    period = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            # use_text_flow keeps long seller names from bleeding into the No. Faktur column
            text = page.extract_text(use_text_flow=True) or ""
            for line in text.splitlines():
                line = line.strip()
                if not period and line.startswith("Dari "):
                    period = line
                if SKIP_RE.match(line):
                    continue
                m = LINE_RE_STRICT.match(line) or LINE_RE_LOOSE.match(line)
                if m:
                    seller = (m.group(1) or "").strip() or NO_SELLER
                    invoices.setdefault((seller, m.group(2)), m.group(3))
    return invoices, period


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")


def build_workbook(invoices, period, source_name, out_path):
    wb = Workbook()

    # --- Sheet Detail: satu baris per faktur unik ---
    det = wb.active
    det.title = "Detail Faktur"
    det.append(["No", "Nama Penjual", "No. Faktur", "Tgl Faktur"])
    style_header(det, 1, 4)
    for i, ((seller, inv), tgl) in enumerate(invoices.items(), start=1):
        det.append([i, seller, inv, tgl])
        for col in range(1, 5):
            c = det.cell(row=i + 1, column=col)
            c.font, c.border = FONT, BORDER
    last_det = len(invoices) + 1
    det.freeze_panes = "A2"
    det.auto_filter.ref = f"A1:D{last_det}"
    for col, width in zip("ABCD", (6, 28, 22, 14)):
        det.column_dimensions[col].width = width

    # --- Sheet Rekap: jumlah faktur per penjual (formula ke sheet Detail) ---
    rek = wb.create_sheet("Rekap", 0)
    rek["A1"] = "Rekap Jumlah Faktur per Penjual"
    rek["A1"].font = Font(name="Arial", size=13, bold=True)
    rek["A2"] = f"{period}  |  Sumber: {source_name}"
    rek["A2"].font = Font(name="Arial", size=9, italic=True, color="595959")

    header_row = 4
    for col, title in enumerate(["No", "Nama Penjual", "Jumlah Faktur", "Daftar No. Faktur"], start=1):
        rek.cell(row=header_row, column=col, value=title)
    style_header(rek, header_row, 4)

    by_seller = {}
    for seller, inv in invoices:
        by_seller.setdefault(seller, []).append(inv)

    row = header_row
    for i, (seller, invs) in enumerate(by_seller.items(), start=1):
        row += 1
        rek.cell(row=row, column=1, value=i)
        rek.cell(row=row, column=2, value=seller)
        rek.cell(row=row, column=3,
                 value=f"=COUNTIF('Detail Faktur'!$B$2:$B${last_det},B{row})")
        rek.cell(row=row, column=4, value=", ".join(sorted(invs)))
        for col in range(1, 5):
            c = rek.cell(row=row, column=col)
            c.font, c.border = FONT, BORDER
        rek.cell(row=row, column=4).alignment = Alignment(wrap_text=True, vertical="top")

    total_row = row + 1
    rek.cell(row=total_row, column=2, value="TOTAL")
    rek.cell(row=total_row, column=3, value=f"=SUM(C{header_row + 1}:C{row})")
    for col in range(1, 5):
        c = rek.cell(row=total_row, column=col)
        c.font, c.fill, c.border = BOLD, TOTAL_FILL, BORDER

    rek.freeze_panes = f"A{header_row + 1}"
    for col, width in zip("ABCD", (6, 28, 15, 90)):
        rek.column_dimensions[col].width = width

    wb.save(out_path)
    return by_seller


def main():
    here = Path(__file__).resolve().parent
    pdfs = [Path(a) for a in sys.argv[1:]] or sorted(here.glob("*.pdf"))
    if not pdfs:
        print("Tidak ada file PDF ditemukan.")
        return

    for pdf_path in pdfs:
        invoices, period = extract_invoices(pdf_path)
        if not invoices:
            print(f"[LEWATI] {pdf_path.name}: tidak ada No. Faktur ditemukan.")
            continue
        out_path = pdf_path.with_name(f"Rekap Faktur - {pdf_path.stem}.xlsx")
        by_seller = build_workbook(invoices, period, pdf_path.name, out_path)

        print(f"\n{pdf_path.name}")
        for seller, invs in by_seller.items():
            print(f"  {seller:<28} {len(invs):>3}")
        print(f"  {'TOTAL':<28} {len(invoices):>3}")
        print(f"  -> {out_path.name}")


if __name__ == "__main__":
    main()
